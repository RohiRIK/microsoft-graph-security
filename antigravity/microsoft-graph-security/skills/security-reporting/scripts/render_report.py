#!/usr/bin/env python3
"""Render a report spec (JSON) into a plain-language Excel workbook and/or a self-contained HTML page.

One spec, two outputs, same pages. The spec holds only what a reader should see; every number in it
must already be computed from source data by the caller (see ../WritingRules.md).

    python3 render_report.py spec.json --xlsx out.xlsx --html out.html
    python3 render_report.py --selftest

Page types (``pages[].type``): summary, table, flow, steps, decisions. See ../Layouts.md.
Requires openpyxl for --xlsx (``pip install openpyxl``); HTML needs nothing beyond the standard library.
"""

from __future__ import annotations

import argparse
import html
import json
import sys
import tempfile
from pathlib import Path

# ── spec validation ──────────────────────────────────────────────────────────

PAGE_TYPES = {"summary", "table", "flow", "steps", "decisions"}


class SpecError(ValueError):
    pass


def validate(spec: dict) -> None:
    for key in ("title", "pages"):
        if key not in spec:
            raise SpecError(f"spec is missing '{key}'")
    if not spec["pages"]:
        raise SpecError("spec has no pages")
    for i, page in enumerate(spec["pages"], start=1):
        kind = page.get("type")
        if kind not in PAGE_TYPES:
            raise SpecError(f"page {i}: type must be one of {sorted(PAGE_TYPES)}, got {kind!r}")
        if not page.get("name"):
            raise SpecError(f"page {i}: needs a 'name' (the tab title)")
        if kind in ("table", "steps", "decisions"):
            headers, rows = page.get("headers"), page.get("rows")
            if not headers or rows is None:
                raise SpecError(f"page {i} ({page['name']}): a {kind} page needs 'headers' and 'rows'")
            for r, row in enumerate(rows, start=1):
                if len(row) != len(headers):
                    raise SpecError(
                        f"page {i} ({page['name']}) row {r}: {len(row)} cells for {len(headers)} headers"
                    )
        if kind == "flow" and not page.get("lines"):
            raise SpecError(f"page {i} ({page['name']}): a flow page needs 'lines'")


def tab_name(i: int, page: dict) -> str:
    """Excel tab: numbered, at most 31 characters, no characters Excel forbids."""
    name = f"{i}. {page['name']}" if page.get("numbered", True) else page["name"]
    for bad in '[]:*?/\\':
        name = name.replace(bad, "-")
    return name[:31]


# ── Excel ────────────────────────────────────────────────────────────────────


def render_xlsx(spec: dict, out: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    font = spec.get("font", "Arial")
    body, bold = Font(name=font, size=11), Font(name=font, size=11, bold=True)
    title_f = Font(name=font, size=18, bold=True, color="1F4E78")
    sub_f = Font(name=font, size=11, italic=True, color="595959")
    head_f = Font(name=font, size=11, bold=True, color="FFFFFF")
    big_f = Font(name=font, size=28, bold=True, color="1F4E78")
    mono, mono_b = Font(name="Courier New", size=11), Font(name="Courier New", size=11, bold=True, color="1F4E78")
    head = PatternFill("solid", fgColor="1F4E78")
    before, after, note = (PatternFill("solid", fgColor=c) for c in ("FCE4D6", "E2EFDA", "FFF2CC"))
    band = PatternFill("solid", fgColor="F2F2F2")
    side = Side(style="thin", color="D9D9D9")
    box = Border(left=side, right=side, top=side, bottom=side)
    wrap = Alignment(wrap_text=True, vertical="top")
    center = Alignment(wrap_text=True, vertical="center", horizontal="center")

    wb = Workbook()
    wb.remove(wb.active)

    def header(ws, row, cells):
        for c, h in enumerate(cells, start=1):
            cell = ws.cell(row=row, column=c, value=h)
            cell.font, cell.fill, cell.border = head_f, head, box
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[row].height = 30

    def grid(ws, top, headers, rows, widths, answer_col=None, banded=True):
        header(ws, top, headers)
        for r, row in enumerate(rows, start=top + 1):
            for c, value in enumerate(row, start=1):
                cell = ws.cell(row=r, column=c, value=value)
                cell.font, cell.border, cell.alignment = body, box, wrap
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    cell.number_format = "#,##0" if isinstance(value, int) else "#,##0.0"
                if answer_col is not None and c == answer_col:
                    cell.fill = note
                elif banded and (r - top) % 2 == 0:
                    cell.fill = band
        for c, w in enumerate(widths or [], start=1):
            ws.column_dimensions[chr(64 + c)].width = w
        if rows:
            ws.freeze_panes = ws.cell(row=top + 1, column=2)
            ws.auto_filter.ref = f"A{top}:{chr(64 + len(headers))}{top + len(rows)}"

    def heading(ws, page, col=1):
        ws.cell(row=1, column=col, value=page.get("title", page["name"])).font = title_f
        if page.get("subtitle"):
            ws.cell(row=2, column=col, value=page["subtitle"]).font = sub_f

    for i, page in enumerate(spec["pages"], start=1):
        ws = wb.create_sheet(tab_name(i, page))
        ws.sheet_view.showGridLines = False
        kind = page["type"]

        if kind == "summary":
            ws.column_dimensions["A"].width = 3
            ws.column_dimensions["B"].width = 34
            ws.column_dimensions["C"].width = 8
            ws.column_dimensions["D"].width = 62
            ws["B2"] = spec["title"] if i == 1 else page.get("title", page["name"])
            ws["B2"].font = title_f
            if spec.get("subtitle") and i == 1:
                ws["B3"] = spec["subtitle"]
                ws["B3"].font = sub_f
            row = 5
            if page.get("before_after"):
                ba = page["before_after"]
                for ref, val, fill, f in (
                    (f"B{row}", ba["before_label"], before, bold), (f"D{row}", ba["after_label"], after, bold),
                    (f"B{row + 1}", ba["before_value"], before, big_f), (f"D{row + 1}", ba["after_value"], after, big_f),
                    (f"B{row + 2}", ba.get("before_caption", ""), before, body),
                    (f"D{row + 2}", ba.get("after_caption", ""), after, body),
                ):
                    ws[ref] = val
                    ws[ref].font, ws[ref].fill, ws[ref].alignment = f, fill, center
                ws[f"C{row + 1}"] = "→"
                ws[f"C{row + 1}"].font, ws[f"C{row + 1}"].alignment = big_f, center
                ws.row_dimensions[row + 1].height, ws.row_dimensions[row + 2].height = 48, 34
                row += 4
            for section in page.get("sections", []):
                ws.cell(row=row, column=2, value=section["heading"]).font = bold
                row += 1
                lines = section.get("bullets") or [section.get("text", "")]
                for line in lines:
                    cell = ws.cell(row=row, column=2, value=(f"• {line}" if section.get("bullets") else line))
                    cell.font, cell.alignment = body, wrap
                    if section.get("callout"):
                        cell.fill = note
                    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
                    ws.row_dimensions[row].height = max(18, 16 * (1 + len(str(line)) // 95))
                    row += 1
                row += 1

        elif kind == "flow":
            heading(ws, page, col=2)
            ws.column_dimensions["A"].width = 3
            ws.column_dimensions["B"].width = max(60, min(140, max(len(l) for l in page["lines"]) + 4))
            for r, line in enumerate(page["lines"], start=4):
                cell = ws.cell(row=r, column=2, value=line)
                cell.font = mono_b if r == 4 and page.get("bold_first_line", True) else mono
                ws.row_dimensions[r].height = 17
            if page.get("note"):
                cell = ws.cell(row=5 + len(page["lines"]), column=2, value=page["note"])
                cell.font, cell.fill, cell.alignment = body, note, wrap

        else:  # table, steps, decisions
            heading(ws, page)
            answer = len(page["headers"]) if kind == "decisions" else page.get("answer_column")
            grid(ws, 4, page["headers"], page["rows"], page.get("widths"), answer_col=answer,
                 banded=kind == "table")
            after_row = 6 + len(page["rows"])
            for n in page.get("notes", []):
                cell = ws.cell(row=after_row, column=1, value=n)
                cell.font, cell.alignment = body, wrap
                ws.merge_cells(start_row=after_row, start_column=1, end_row=after_row,
                               end_column=len(page["headers"]))
                ws.row_dimensions[after_row].height = max(18, 16 * (1 + len(n) // 110))
                after_row += 1

    wb.save(out)


# ── HTML ─────────────────────────────────────────────────────────────────────

CSS = """
:root{--ink:#1f2328;--muted:#59636e;--accent:#1f4e78;--line:#d9d9d9;--band:#f6f8fa;--before:#fce4d6;
--after:#e2efda;--note:#fff2cc;--bg:#ffffff}
@media (prefers-color-scheme:dark){:root{--ink:#e6edf3;--muted:#9da7b3;--accent:#79b8ff;--line:#30363d;
--band:#161b22;--before:#4a2a1a;--after:#1f3a24;--note:#3d3413;--bg:#0d1117}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--ink);
font:16px/1.55 -apple-system,Segoe UI,Arial,sans-serif}
main{max-width:1100px;margin:0 auto;padding:24px 16px 64px}
header h1{color:var(--accent);margin:0 0 4px;font-size:1.8rem}.sub{color:var(--muted);font-style:italic;margin:0}
nav{margin:20px 0;display:flex;flex-wrap:wrap;gap:8px}nav a{color:var(--accent);text-decoration:none;
border:1px solid var(--line);border-radius:999px;padding:4px 12px;font-size:.9rem}
section{margin:40px 0}h2{color:var(--accent);border-bottom:2px solid var(--line);padding-bottom:6px}
.ba{display:grid;grid-template-columns:1fr auto 1fr;gap:12px;align-items:center;margin:16px 0 24px}
.ba div{border-radius:10px;padding:16px;text-align:center}.ba .b{background:var(--before)}.ba .a{background:var(--after)}
.ba strong{display:block;font-size:2.6rem;color:var(--accent);line-height:1.1}.arrow{font-size:2rem;color:var(--accent)}
.callout{background:var(--note);border-radius:8px;padding:12px 16px}
.scroll{overflow-x:auto}table{border-collapse:collapse;width:100%;font-size:.95rem}
th{background:var(--accent);color:#fff;text-align:left;padding:8px;vertical-align:bottom}
td{border:1px solid var(--line);padding:8px;vertical-align:top}tbody tr:nth-child(even){background:var(--band)}
td.num{text-align:right;font-variant-numeric:tabular-nums}td.answer{background:var(--note)}
pre.flow{font:14px/1.45 ui-monospace,Menlo,Consolas,monospace;background:var(--band);border:1px solid var(--line);
border-radius:8px;padding:16px;overflow-x:auto}
@media print{nav{display:none}section{break-inside:avoid-page}}
"""


def esc(v) -> str:
    if isinstance(v, int) and not isinstance(v, bool):
        return f"{v:,}"
    return html.escape("" if v is None else str(v))


def render_html(spec: dict, out: Path) -> None:
    parts = [f"<!doctype html><html lang='en'><head><meta charset='utf-8'>"
             f"<meta name='viewport' content='width=device-width,initial-scale=1'>"
             f"<title>{esc(spec.get('short_title', spec['title']))}</title><style>{CSS}</style></head><body><main>",
             f"<header><h1>{esc(spec['title'])}</h1>"
             + (f"<p class='sub'>{esc(spec['subtitle'])}</p>" if spec.get("subtitle") else "") + "</header>",
             "<nav>" + "".join(f"<a href='#p{i}'>{i}. {esc(p['name'])}</a>"
                              for i, p in enumerate(spec["pages"], start=1)) + "</nav>"]
    for i, page in enumerate(spec["pages"], start=1):
        parts.append(f"<section id='p{i}'><h2>{i}. {esc(page.get('title', page['name']))}</h2>")
        if page.get("subtitle"):
            parts.append(f"<p class='sub'>{esc(page['subtitle'])}</p>")
        kind = page["type"]
        if kind == "summary":
            if page.get("before_after"):
                ba = page["before_after"]
                parts.append(
                    "<div class='ba'>"
                    f"<div class='b'>{esc(ba['before_label'])}<strong>{esc(ba['before_value'])}</strong>"
                    f"{esc(ba.get('before_caption', ''))}</div><span class='arrow'>→</span>"
                    f"<div class='a'>{esc(ba['after_label'])}<strong>{esc(ba['after_value'])}</strong>"
                    f"{esc(ba.get('after_caption', ''))}</div></div>")
            for s in page.get("sections", []):
                parts.append(f"<h3>{esc(s['heading'])}</h3>")
                if s.get("bullets"):
                    parts.append("<ul>" + "".join(f"<li>{esc(b)}</li>" for b in s["bullets"]) + "</ul>")
                else:
                    cls = " class='callout'" if s.get("callout") else ""
                    parts.append(f"<p{cls}>{esc(s.get('text', ''))}</p>")
        elif kind == "flow":
            parts.append("<pre class='flow'>" + esc("\n".join(page["lines"])) + "</pre>")
            if page.get("note"):
                parts.append(f"<p class='callout'>{esc(page['note'])}</p>")
        else:
            answer = len(page["headers"]) if kind == "decisions" else page.get("answer_column")
            parts.append("<div class='scroll'><table><thead><tr>"
                         + "".join(f"<th>{esc(h)}</th>" for h in page["headers"]) + "</tr></thead><tbody>")
            for row in page["rows"]:
                cells = []
                for c, v in enumerate(row, start=1):
                    cls = "answer" if c == answer else ("num" if isinstance(v, (int, float)) and not isinstance(v, bool) else "")
                    cells.append(f"<td class='{cls}'>{esc(v)}</td>" if cls else f"<td>{esc(v)}</td>")
                parts.append("<tr>" + "".join(cells) + "</tr>")
            parts.append("</tbody></table></div>")
            for n in page.get("notes", []):
                parts.append(f"<p class='callout'>{esc(n)}</p>")
        parts.append("</section>")
    parts.append("</main></body></html>")
    out.write_text("".join(parts), encoding="utf-8")


# ── entry point ──────────────────────────────────────────────────────────────


def selftest() -> int:
    example = Path(__file__).resolve().parent.parent / "examples" / "merge-plan.spec.json"
    spec = json.loads(example.read_text(encoding="utf-8"))
    validate(spec)
    with tempfile.TemporaryDirectory() as tmp:
        x, h = Path(tmp) / "t.xlsx", Path(tmp) / "t.html"
        render_html(spec, h)
        assert h.read_text(encoding="utf-8").count("<section") == len(spec["pages"])
        try:
            render_xlsx(spec, x)
            from openpyxl import load_workbook

            assert len(load_workbook(x).sheetnames) == len(spec["pages"])
            xlsx = "xlsx ok"
        except ImportError:
            xlsx = "xlsx skipped (pip install openpyxl)"
    try:
        validate({"title": "t", "pages": [{"type": "table", "name": "x", "headers": ["a", "b"], "rows": [[1]]}]})
        raise AssertionError("a short row must be rejected")
    except SpecError:
        pass
    print(f"selftest ok — {len(spec['pages'])} pages, html ok, {xlsx}")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("spec", nargs="?", help="report spec JSON")
    ap.add_argument("--xlsx", type=Path)
    ap.add_argument("--html", type=Path)
    ap.add_argument("--selftest", action="store_true")
    a = ap.parse_args()
    if a.selftest:
        return selftest()
    if not a.spec or not (a.xlsx or a.html):
        ap.error("give a spec and at least one of --xlsx / --html")
    spec = json.loads(Path(a.spec).read_text(encoding="utf-8"))
    try:
        validate(spec)
    except SpecError as e:
        print(f"spec error: {e}", file=sys.stderr)
        return 2
    if a.xlsx:
        render_xlsx(spec, a.xlsx)
        print(f"wrote {a.xlsx}")
    if a.html:
        render_html(spec, a.html)
        print(f"wrote {a.html}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
